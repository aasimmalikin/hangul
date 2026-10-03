"""create_file: turn content into a real document the user can download.

Formats: PDF, Word (.docx), PowerPoint (.pptx), Excel (.xlsx), CSV, Markdown
and plain text. Documents take simple Markdown (# headings, - bullets,
1. lists, **bold**, | tables |); a deck starts a new slide at every `#`
heading; spreadsheets take ``sheets`` (rows, first row = header).

Files land in the user's own folder (``data/sessions/<user>``) under a fresh
name -- an existing file is never overwritten, which is why this is tier SAFE
and there is no tool that edits or deletes a file. The chat shows a download
card (``GET /files/<name>``). Plan-gated: Plus and up (billing/plans.py).
"""

import asyncio
import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

from harness.tools.base import Tool, ToolOutput

SESSIONS = Path("data/sessions")
MAX_CHARS = 200_000
MAX_ROWS = 20_000
FORMATS = {"pdf": "pdf", "docx": "docx", "word": "docx", "pptx": "pptx", "slides": "pptx", "powerpoint": "pptx",
           "xlsx": "xlsx", "excel": "xlsx", "csv": "csv", "md": "md", "markdown": "md", "txt": "txt", "text": "txt"}
FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/usr/local/share/fonts")


def user_folder(user_id: str) -> Path:
    d = SESSIONS / "".join(c for c in user_id if c.isalnum() or c in "-_")
    d.mkdir(parents=True, exist_ok=True)
    return d


def _slug(title: str, ext: str) -> str:
    # a title that already ends in the extension ("notes.txt") keeps its name, not "notes-txt"
    if title.lower().endswith(f".{ext}"):
        title = title[: -len(ext) - 1]
    return re.sub(r"[^A-Za-z0-9]+", "-", title).strip("-").lower()[:60] or "document"


def fresh_name(folder: Path, title: str, ext: str) -> str:
    slug = _slug(title, ext)
    name, n = f"{slug}.{ext}", 2
    while (folder / name).exists():
        name, n = f"{slug}-{n}.{ext}", n + 1
    return name


# ------------------------------------------------------------------ parsing

@dataclass
class Block:
    kind: str                 # h | p | li | ol | table | hr
    text: str = ""
    level: int = 0
    rows: list | None = None


def parse_markdown(md: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    table: list[list[str]] = []

    def flush_para():
        if para:
            blocks.append(Block("p", " ".join(para).strip()))
            para.clear()

    def flush_table():
        if table:
            blocks.append(Block("table", rows=[r for r in table if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]))
            table.clear()

    for raw in md.splitlines():
        line = raw.rstrip()
        s = line.strip()
        if s.startswith("|") and s.endswith("|") and len(s) > 1:
            flush_para()
            table.append([c.strip() for c in s.strip("|").split("|")])
            continue
        flush_table()
        if not s:
            flush_para()
        elif m := re.match(r"^(#{1,3})\s+(.*)", s):
            flush_para()
            blocks.append(Block("h", m.group(2).strip(), level=len(m.group(1))))
        elif re.fullmatch(r"(-{3,}|\*{3,})", s):
            flush_para()
            blocks.append(Block("hr"))
        elif m := re.match(r"^[-*•]\s+(.*)", s):
            flush_para()
            blocks.append(Block("li", m.group(1)))
        elif m := re.match(r"^\d+[.)]\s+(.*)", s):
            flush_para()
            blocks.append(Block("ol", m.group(1)))
        else:
            para.append(s)
    flush_para()
    flush_table()
    return blocks


def plain(text: str) -> str:
    return re.sub(r"(\*\*|__|`)", "", text)


def runs(text: str) -> list[tuple[str, bool]]:
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]."""
    out, bold = [], False
    for part in re.split(r"\*\*|__", text):
        if part:
            out.append((part.replace("`", ""), bold))
        bold = not bold
    return out


def _cell(v):
    """Spreadsheet cells: numbers stay numbers ('1,240' -> 1240)."""
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v).strip()
    t = s.replace(",", "")
    if re.fullmatch(r"-?\d+", t):
        return int(t)
    if re.fullmatch(r"-?\d*\.\d+", t):
        return float(t)
    return s


# ---------------------------------------------------------------- writers

def write_docx(title: str, blocks: list[Block]) -> bytes:
    from docx import Document
    doc = Document()
    doc.core_properties.title = title
    if not blocks or blocks[0].kind != "h":
        doc.add_heading(title, level=0)
    for i, b in enumerate(blocks):
        if b.kind == "h":
            # a leading heading is the document title
            doc.add_heading(plain(b.text), level=0 if i == 0 else min(b.level, 3))
        elif b.kind in ("p", "li", "ol"):
            para = doc.add_paragraph(style={"p": None, "li": "List Bullet", "ol": "List Number"}[b.kind])
            for text, bold in runs(b.text):
                para.add_run(text).bold = bold
        elif b.kind == "table" and b.rows:
            width = max(len(r) for r in b.rows)
            t = doc.add_table(rows=0, cols=width)
            t.style = "Table Grid"
            for ri, row in enumerate(b.rows):
                cells = t.add_row().cells
                for ci in range(width):
                    cells[ci].text = plain(row[ci]) if ci < len(row) else ""
                    if ri == 0:
                        for p in cells[ci].paragraphs:
                            for r in p.runs:
                                r.bold = True
        elif b.kind == "hr":
            doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _font_files() -> tuple[str, str] | None:
    for d in FONT_DIRS:
        reg, bold = Path(d) / "DejaVuSans.ttf", Path(d) / "DejaVuSans-Bold.ttf"
        if reg.is_file():
            return str(reg), str(bold if bold.is_file() else reg)
    return None


def write_pdf(title: str, blocks: list[Block]) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_title(title)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    fonts = _font_files()
    if fonts:
        pdf.add_font("Body", "", fonts[0])
        pdf.add_font("Body", "B", fonts[1])
        family = "Body"
        clean = lambda s: s  # noqa: E731
    else:
        # no Unicode font installed: core font, with characters it lacks replaced
        family = "Helvetica"
        clean = lambda s: s.encode("latin-1", "replace").decode("latin-1")  # noqa: E731
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def para(text: str, size: float = 11, style: str = "", indent: str = ""):
        pdf.set_font(family, style, size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(width, size * 0.55, clean(indent + text), markdown=not style, new_x="LMARGIN", new_y="NEXT")

    if not blocks or blocks[0].kind != "h":
        para(title, 20, "B")
        pdf.ln(3)
    n = 0
    for b in blocks:
        if b.kind == "h":
            pdf.ln(2)
            para(plain(b.text), {1: 18, 2: 14, 3: 12}.get(b.level, 12), "B")
            pdf.ln(1)
        elif b.kind == "p":
            para(b.text)
            pdf.ln(2)
        elif b.kind in ("li", "ol"):
            n = n + 1 if b.kind == "ol" else n
            para(b.text, indent="•  " if b.kind == "li" else f"{n}.  ")
        elif b.kind == "table" and b.rows:
            pdf.set_font(family, "", 10)
            cols = max(len(r) for r in b.rows)
            with pdf.table(text_align="LEFT", line_height=6) as t:
                for row in b.rows:
                    r = t.row()
                    for ci in range(cols):
                        r.cell(clean(plain(row[ci])) if ci < len(row) else "")
            pdf.ln(3)
        elif b.kind == "hr":
            pdf.add_page()
        if b.kind != "ol":
            n = 0                                   # a numbered list restarts after anything else
    return bytes(pdf.output())


def write_pptx(title: str, blocks: list[Block]) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    slides: list[tuple[str, list[Block]]] = []
    for b in blocks:
        if b.kind == "h" or not slides:
            slides.append((plain(b.text) if b.kind == "h" else title, []))
            if b.kind == "h":
                continue
        if b.kind != "hr":
            slides[-1][1].append(b)
    if not slides:
        slides = [(title, [])]
    for i, (heading, body) in enumerate(slides):
        if i == 0 and not body:
            s = prs.slides.add_slide(prs.slide_layouts[0])          # title slide
            s.shapes.title.text = heading
            continue
        s = prs.slides.add_slide(prs.slide_layouts[1])              # title + content
        s.shapes.title.text = heading
        tf = s.placeholders[1].text_frame
        tf.clear()
        first = True
        for b in body:
            lines = [" · ".join(plain(c) for c in r) for r in (b.rows or [])] if b.kind == "table" else [plain(b.text)]
            for line in lines:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.text = line
                p.font.size = Pt(18 if b.kind in ("li", "ol", "table") else 20)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def write_xlsx(sheets: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(sheets or [{"name": "Sheet1", "rows": []}]):
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", " ", str(sh.get("name") or f"Sheet{i + 1}"))[:31])
        for r in (sh.get("rows") or [])[:MAX_ROWS]:
            ws.append([_cell(v) for v in (r if isinstance(r, list) else [r])])
        if ws.max_row:
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_csv(sheets: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in ((sheets or [{}])[0].get("rows") or [])[:MAX_ROWS]:
        w.writerow(r if isinstance(r, list) else [r])
    return buf.getvalue().encode("utf-8-sig")       # BOM: Excel opens UTF-8 correctly


def build(fmt: str, title: str, content: str, sheets: list[dict] | None) -> bytes:
    if fmt in ("xlsx", "csv"):
        if not sheets and content:
            sheets = [{"name": title[:31], "rows": [row for b in parse_markdown(content) if b.kind == "table" for row in b.rows]}]
        return write_xlsx(sheets or []) if fmt == "xlsx" else write_csv(sheets or [])
    if fmt in ("md", "txt"):
        return (content if fmt == "md" else "\n".join(plain(re.sub(r"^\s*#{1,6}\s*", "", line))
                                                     for line in content.splitlines())).encode()
    blocks = parse_markdown(content)
    return {"docx": write_docx, "pdf": write_pdf, "pptx": write_pptx}[fmt](title, blocks)


def make_create_file_tool(user_id: str) -> Tool:
    async def create_file(title: str, format: str, content: str = "", sheets: list[dict] | None = None) -> ToolOutput | str:  # noqa: A002
        fmt = FORMATS.get(format.strip().lower().lstrip("."))
        if not fmt:
            return f"Unsupported format {format!r}. Use one of: pdf, docx, pptx, xlsx, csv, md, txt."
        if len(content) > MAX_CHARS:
            return f"That is too long for one file ({len(content)} characters; max {MAX_CHARS})."
        if fmt in ("xlsx", "csv") and not sheets and "|" not in content:
            return "A spreadsheet needs `sheets` (each with `rows`, the first row being the header) or a Markdown table."
        if fmt not in ("xlsx", "csv") and not content.strip():
            return "Give the document's content (Markdown)."
        try:
            data = await asyncio.to_thread(build, fmt, title.strip() or "Document", content, sheets)
        except Exception as e:  # noqa: BLE001 - a bad table must not crash the run
            return f"Could not create the file: {type(e).__name__}: {e}"
        folder = user_folder(user_id)
        name = fresh_name(folder, title, fmt)
        await asyncio.to_thread((folder / name).write_bytes, data)
        size_kb = max(1, round(len(data) / 1024))
        renamed = "" if name == f"{_slug(title, fmt)}.{fmt}" else (
            f" A file called {_slug(title, fmt)}.{fmt} already existed and is never overwritten, so this one is "
            f"{name}; tell the user that name.")
        return ToolOutput(f"Created {name} ({size_kb} KB).{renamed} The user can download it from the card shown "
                          "in the chat.",
                          {"kind": "file", "name": name, "format": fmt, "size": len(data), "title": title})

    return Tool(
        name="create_file",
        description=(
            "Create or save a new file in the user's folder (they can download it): pdf, docx (Word), pptx (slides), "
            "xlsx (Excel), csv, md or txt -- this is also how to save text to a file such as notes.txt. "
            "For documents and slides pass `content` as Markdown (# headings, - bullets, **bold**, | tables |; "
            "in slides every # heading starts a new slide). For xlsx/csv pass `sheets`: [{name, rows}] with the "
            "first row as the header (numbers as numbers). Use when the user asks for a file, report, document, "
            "spreadsheet, slides or 'something I can download/send/print'. Write the full content yourself."),
        parameter={
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Human title; also becomes the file name."},
                "format": {"type": "string", "enum": ["pdf", "docx", "pptx", "xlsx", "csv", "md", "txt"]},
                "content": {"type": "string", "description": "Markdown body (documents, slides, md, txt)."},
                "sheets": {"type": "array", "items": {"type": "object", "properties": {
                    "name": {"type": "string"},
                    "rows": {"type": "array", "items": {"type": "array", "items": {}}}}}},
            },
            "required": ["title", "format"],
        },
        handler=create_file,
    )
