"""Turn an uploaded file's bytes into clean text chunks, ready to embed."""

from io import BytesIO

MAX_BYTES = 10*1024*1024
TEXT_TYPES = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
AUDIO_TYPES = {".mp3", ".m4a", ".wav", ".webm", ".ogg", ".oga", ".mpga"}
ALLOWED = TEXT_TYPES | IMAGE_TYPES | AUDIO_TYPES
MAX_SHEET_ROWS = 5000

class UploadError(Exception):
    """Raised for unsupported type or oversize upload; mapped to HTTP"""

def ext_of(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

def is_image(filename: str) -> bool:
    return ext_of(filename) in IMAGE_TYPES

def is_audio(filename: str) -> bool:
    return ext_of(filename) in AUDIO_TYPES

def extract_text(filename:str, data: bytes)->str:
    """Text of a document. Images are not handled here: they need a model call
    (``media.vision``), which /upload makes and then indexes like this text."""
    ext = ext_of(filename)
    if ext not in TEXT_TYPES:
        raise UploadError(f"Unsupported file type '{ext}'. Allowed: {sorted(ALLOWED)}")
    if len(data) > MAX_BYTES:
        raise UploadError(f"File too large ({len(data)} bytes). Max {MAX_BYTES // (1024*1024)} MB.")

    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    if ext == ".csv":
        return _csv_text(data.decode("utf-8-sig", errors="replace"))
    if ext == ".docx":
        return _docx_text(data)
    if ext == ".xlsx":
        return _xlsx_text(data)

    from pypdf import PdfReader
    reader = PdfReader(BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _rows_text(rows, title: str = "") -> str:
    """Rows as 'Header: value; Header: value' lines, so each chunk keeps its
    column names and search finds 'the March rent row', not bare numbers."""
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    head, body = rows[0], rows[1:MAX_SHEET_ROWS]
    lines = [f"{title}" if title else "", "Columns: " + ", ".join(h or f"col{i+1}" for i, h in enumerate(head))]
    for r in body:
        lines.append("; ".join(f"{(head[i] if i < len(head) and head[i] else f'col{i+1}')}: {v}"
                               for i, v in enumerate(r) if v))
    return "\n".join(x for x in lines if x)


def _csv_text(text: str) -> str:
    import csv
    from io import StringIO
    return _rows_text(csv.reader(StringIO(text)))


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        raise UploadError("Word files are not supported on this server (python-docx missing).")
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        parts.append(_rows_text([[c.text for c in row.cells] for row in t.rows]))
    return "\n\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise UploadError("Excel files are not supported on this server (openpyxl missing).")
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    return "\n\n".join(_rows_text(ws.iter_rows(values_only=True), title=f"Sheet: {ws.title}") for ws in wb.worksheets)


def chunk_text(text: str, size: int = 120, overlap: int = 20) -> list[str]:
    words = text.split()
    if not words:
        return []
    chunks, start = [], 0
    step = max(1, size - overlap)
    while start < len(words):
        chunks.append(" ".join(words[start:start + size]))
        start += step
    return chunks
