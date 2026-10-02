"""The everyday-assistant tools (tools/builtin/daily.py): reminders, lists,
notes, link reader, weather, converter, world clock, image viewer, the wider
upload types, and reminder delivery. Fakes only: no DB, no network, no model.
"""
import asyncio
import io
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from harness import notify, scheduler
from harness.api.auth import get_current_user
from harness.api.routes import personal as personal_route
from harness.db import personal as db
from harness.retrieval import upload_ingest
from harness.tools.base import ToolOutput
from harness.tools.builtin import convert, personal, view_image, weather, web_reader
from harness.tools.builtin.daily import DAILY_TOOL_NAMES, build_daily_tools


def run(coro):
    return asyncio.run(coro)


# ------------------------------------------------------------ fake store

@pytest.fixture
def store(monkeypatch):
    """In-memory stand-in for harness.db.personal, isolated per user like the SQL."""
    st = {"rem": [], "todo": [], "notes": [], "seq": 0}

    def nid():
        st["seq"] += 1
        return st["seq"]

    def add_reminder(uid, text, due):
        r = SimpleNamespace(id=nid(), user=uid, text=text, due_at=due.astimezone(UTC), status="pending", sent_at=None)
        st["rem"].append(r)
        return db.ReminderOut(r.id, r.text, r.due_at.isoformat(), r.status, None)

    def list_reminders(uid, statuses=("pending", "sent"), limit=50):
        return [db.ReminderOut(r.id, r.text, r.due_at.isoformat(), r.status, None)
                for r in st["rem"] if r.user == uid and r.status in statuses]

    def set_reminder_status(uid, rid, status):
        for r in st["rem"]:
            if r.id == rid and r.user == uid:
                r.status = status
                return db.ReminderOut(r.id, r.text, r.due_at.isoformat(), r.status, None)
        return None

    def add_todos(uid, list_name, items):
        out = []
        for i in items:
            t = db.TodoOut(nid(), db.norm_list(list_name), i, False)
            st["todo"].append((uid, t))
            out.append(t)
        return out

    def list_todos(uid, list_name=None, include_done=False):
        return [t for u, t in st["todo"] if u == uid and (include_done or not t.done)
                and (list_name is None or t.list_name.lower() == db.norm_list(list_name).lower())]

    def find_todos(uid, text, list_name=None):
        return [t for t in list_todos(uid, list_name) if text.lower() in t.text.lower()]

    def set_todo_done(uid, iid, done):
        for u, t in st["todo"]:
            if u == uid and t.id == iid:
                t.done = done
                return t
        return None

    def delete_todo(uid, iid):
        before = len(st["todo"])
        st["todo"] = [(u, t) for u, t in st["todo"] if not (u == uid and t.id == iid)]
        return len(st["todo"]) < before

    def add_note(uid, text):
        n = db.NoteOut(nid(), text, datetime.now(UTC).isoformat())
        st["notes"].append((uid, n))
        return n

    def search_notes(uid, query="", limit=20):
        words = [w.lower() for w in query.split() if len(w) > 2]
        return [n for u, n in st["notes"] if u == uid and (not words or any(w in n.text.lower() for w in words))]

    for name, fn in dict(add_reminder=add_reminder, list_reminders=list_reminders,
                         set_reminder_status=set_reminder_status, add_todos=add_todos, list_todos=list_todos,
                         find_todos=find_todos, set_todo_done=set_todo_done, delete_todo=delete_todo,
                         list_names=lambda uid: sorted({t.list_name for u, t in st["todo"] if u == uid}),
                         add_note=add_note, search_notes=search_notes,
                         delete_note=lambda uid, i: False).items():
        monkeypatch.setattr(db, name, fn)
    return st


def test_daily_toolset_is_complete():
    assert [t.name for t in build_daily_tools("7", "Asia/Kolkata")] == list(DAILY_TOOL_NAMES)


# ------------------------------------------------------------ reminders

def test_reminder_local_time_is_read_in_the_users_timezone(store):
    tool = personal.make_reminders_tool("7", "Asia/Kolkata")
    when = (datetime.now(UTC) + timedelta(days=1)).astimezone().strftime("%Y-%m-%d") + "T19:00"
    out = run(tool.handler(action="add", text="Call mom", when=when))
    assert isinstance(out, ToolOutput) and out.ui["kind"] == "reminder"
    due = store["rem"][0].due_at
    assert due.astimezone(personal.ZoneInfo("Asia/Kolkata")).strftime("%H:%M") == "19:00"
    assert due.strftime("%H:%M") == "13:30"                       # stored as UTC


def test_reminder_in_minutes_and_past_times(store):
    tool = personal.make_reminders_tool("7", "UTC")
    out = run(tool.handler(action="add", text="Stretch", in_minutes=30))
    assert "Reminder #" in out.text
    assert "already passed" in run(tool.handler(action="add", text="x", when="2020-01-01T10:00"))
    assert "Give a time" in run(tool.handler(action="add", text="x"))


def test_reminders_list_and_cancel_are_per_user(store):
    mine = personal.make_reminders_tool("7", "UTC")
    theirs = personal.make_reminders_tool("8", "UTC")
    run(mine.handler(action="add", text="Mine", in_minutes=10))
    assert "no upcoming" in run(theirs.handler(action="list"))
    rid = store["rem"][0].id
    assert "No reminder" in run(theirs.handler(action="cancel", reminder_id=rid))
    assert "Cancelled" in run(mine.handler(action="cancel", reminder_id=rid))


# ---------------------------------------------------------------- lists

def test_lists_add_show_done(store):
    tool = personal.make_lists_tool("7")
    out = run(tool.handler(action="add", list="shopping", items=["milk", "eggs"]))
    assert out.ui["kind"] == "todo_list" and out.ui["list"] == "Shopping"
    assert [i["text"] for i in out.ui["items"]] == ["milk", "eggs"]
    out = run(tool.handler(action="done", list="Shopping", item="egg"))
    assert "Ticked off: eggs" in out.text
    assert [i["text"] for i in out.ui["items"]] == ["milk"]


def test_lists_ambiguous_item_asks(store):
    tool = personal.make_lists_tool("7")
    run(tool.handler(action="add", items=["call bank", "call plumber"]))
    assert "Several items match" in run(tool.handler(action="done", item="call"))


# ---------------------------------------------------------------- notes

def test_notes_add_and_search(store):
    tool = personal.make_notes_tool("7")
    run(tool.handler(action="add", text="Car service due in March"))
    run(tool.handler(action="add", text="Wifi password is on the router"))
    out = run(tool.handler(action="search", query="car service"))
    assert "Car service" in out.text and "Wifi" not in out.text
    assert "No notes" in run(personal.make_notes_tool("8").handler(action="search", query="car"))


# ------------------------------------------------------------ converter

def test_unit_conversions():
    assert convert.convert_units(1, "mile", "km") == pytest.approx(1.609344)
    assert convert.convert_units(100, "C", "F") == pytest.approx(212)
    assert convert.convert_units(5, "ft", "cm") == pytest.approx(152.4)
    with pytest.raises(ValueError):
        convert.convert_units(1, "kg", "km")


def test_currency_conversion_uses_the_rate_service(monkeypatch):
    class FakeClient:
        def __init__(self, *a, **k): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False
        async def get(self, url, params):
            assert params == {"amount": 100.0, "from": "USD", "to": "INR"}
            return SimpleNamespace(status_code=200, raise_for_status=lambda: None,
                                   json=lambda: {"date": "2026-10-01", "rates": {"INR": 8350.5}})
    monkeypatch.setattr(convert.httpx, "AsyncClient", FakeClient)
    out = run(convert.convert(100, "usd", "inr"))
    assert "8,350.5 INR" in out.text and out.ui["kind"] == "conversion"


def test_world_clock_at_a_local_time():
    tool = convert.make_world_clock_tool("Asia/Kolkata")
    out = run(tool.handler(places=["London", "Asia/Tokyo"], at="2026-01-15T18:00"))
    rows = {r["place"]: r["time"] for r in out.ui["rows"]}
    assert rows == {"London": "12:30", "Asia/Tokyo": "21:30"}
    assert "Unknown place" in run(tool.handler(places=["Atlantis"]))


# --------------------------------------------------------------- weather

def test_weather_card(monkeypatch):
    async def fake_get(client, url, params):
        if url == weather.GEO:
            return {"results": [{"name": "Pune", "admin1": "Maharashtra", "country": "India",
                                 "country_code": "IN", "latitude": 18.5, "longitude": 73.8}]}
        return {"timezone": "Asia/Kolkata",
                "current": {"temperature_2m": 27.4, "apparent_temperature": 29, "relative_humidity_2m": 70,
                            "weather_code": 61, "wind_speed_10m": 12},
                "daily": {"time": ["2026-10-02", "2026-10-03"], "weather_code": [61, 0],
                          "temperature_2m_max": [29, 31], "temperature_2m_min": [22, 23],
                          "precipitation_probability_max": [80, 5]}}
    monkeypatch.setattr(weather, "_get", fake_get)
    out = run(weather.weather("Pune", days=2))
    assert out.ui["kind"] == "weather" and out.ui["place"].startswith("Pune")
    assert out.ui["current"]["label"] == "Light rain"
    assert "rain chance 80%" in out.text


# ------------------------------------------------------------ link reader

def test_html_to_text_drops_page_chrome():
    title, text = web_reader.html_to_text(
        "<html><head><title>Rent rules</title><script>evil()</script></head><body><nav>Home | About</nav>"
        "<h1>Deposits</h1><p>Max two months.</p><ul><li>Refund in 30 days</li></ul><footer>(c)</footer></body></html>")
    assert title == "Rent rules"
    assert "## Deposits" in text and "Max two months." in text and "- Refund in 30 days" in text
    assert "evil" not in text and "Home | About" not in text and "(c)" not in text


@pytest.mark.parametrize("url", ["http://127.0.0.1/admin", "http://localhost:8000/healthz", "http://10.0.0.5/",
                                 "http://169.254.169.254/latest/meta-data", "file:///etc/passwd",
                                 "http://user:pw@example.com/"])
def test_link_reader_refuses_private_and_odd_urls(url):
    out = run(web_reader.read_webpage(url))
    assert out.startswith("Could not read that link")


# --------------------------------------------------------------- images

def test_view_image_stays_in_the_users_folder(tmp_path, monkeypatch):
    monkeypatch.setattr(view_image, "SESSIONS", tmp_path)
    (tmp_path / "7").mkdir()
    (tmp_path / "7" / "bill.jpg").write_bytes(b"\xff\xd8fake")
    (tmp_path / "8").mkdir()
    (tmp_path / "8" / "secret.png").write_bytes(b"x")
    seen = {}

    async def fake_ask(data, mime, question, **kw):
        seen.update(data=data, mime=mime, user=kw.get("user_id"))
        return "Total: 1,240 INR, due 5 Oct"
    monkeypatch.setattr(view_image, "ask_image", fake_ask)
    tool = view_image.make_view_image_tool("7")
    out = run(tool.handler(filename="../8/secret.png"))
    assert "No image named 'secret.png'" in out and "bill.jpg" in out      # path forced to own folder
    out = run(tool.handler(filename="bill.jpg", question="total?"))
    assert "1,240 INR" in out and seen == {"data": b"\xff\xd8fake", "mime": "image/jpeg", "user": "7"}


# --------------------------------------------------------------- uploads

def test_csv_rows_keep_their_column_names():
    text = upload_ingest.extract_text("budget.csv", b"Month,Rent,Food\nMarch,15000,6200\n")
    assert "Columns: Month, Rent, Food" in text
    assert "Month: March; Rent: 15000; Food: 6200" in text


def test_docx_and_xlsx_are_read():
    from docx import Document
    from openpyxl import Workbook
    d = Document()
    d.add_paragraph("Lease starts 1 April.")
    buf = io.BytesIO()
    d.save(buf)
    assert "Lease starts 1 April." in upload_ingest.extract_text("lease.docx", buf.getvalue())
    wb = Workbook()
    wb.active.title = "Bills"
    wb.active.append(["Item", "Amount"])
    wb.active.append(["Power", 1240])
    buf = io.BytesIO()
    wb.save(buf)
    text = upload_ingest.extract_text("bills.xlsx", buf.getvalue())
    assert "Sheet: Bills" in text and "Item: Power; Amount: 1240" in text


def test_images_are_allowed_but_not_text_extracted():
    assert upload_ingest.is_image("receipt.JPG")
    with pytest.raises(upload_ingest.UploadError):
        upload_ingest.extract_text("receipt.jpg", b"x")
    with pytest.raises(upload_ingest.UploadError):
        upload_ingest.extract_text("movie.mp4", b"x")


# ------------------------------------------------------------ routes

def test_routes_hide_other_users_items(store):
    app = FastAPI()
    app.include_router(personal_route.router)
    app.dependency_overrides[get_current_user] = lambda: {"user_id": "8"}
    run(personal.make_lists_tool("7").handler(action="add", items=["milk"]))
    item_id = store["todo"][0][1].id
    c = TestClient(app)
    assert c.get("/lists").json() == {"lists": {}}
    assert c.patch(f"/lists/items/{item_id}", json={"done": True}).status_code == 404
    assert c.post("/reminders/999/done").status_code == 404
    assert c.get("/reminders?scope=nope").status_code == 422


# ------------------------------------------------------------ delivery

def test_due_reminders_are_claimed_once_and_emailed(monkeypatch):
    due = SimpleNamespace(id=5, text="Pay rent", due_at=datetime(2026, 10, 1, 3, 30, tzinfo=UTC))
    claimed, emailed, sent = set(), [], []
    monkeypatch.setattr(db, "due_reminders", lambda: [(7, due), (7, due)])   # seen twice (e.g. two workers)
    monkeypatch.setattr(db, "claim_due", lambda rid: not (rid in claimed or claimed.add(rid)))
    monkeypatch.setattr(db, "mark_emailed", lambda rid: emailed.append(rid))
    monkeypatch.setattr(notify, "email_enabled", lambda: True)
    monkeypatch.setattr(notify, "user_email", lambda uid: "a@example.com")

    async def fake_send(to, subject, text, html=None):
        sent.append((to, subject))
        return True
    monkeypatch.setattr(notify, "send_email", fake_send)
    import harness.db.settings as settings_db
    monkeypatch.setattr(settings_db, "get_settings", lambda uid: SimpleNamespace(timezone="Asia/Kolkata"))

    assert run(scheduler.fire_reminders()) == 1
    assert sent == [("a@example.com", "⏰ Reminder: Pay rent")] and emailed == [5]
