"""Phase 3 productivity tools: create_file (PDF/Word/slides/Excel/CSV/MD/TXT),
analyze_data (fixed pandas operations + charts, no code execution), plan
gating (Free gets an upgrade card instead), and GET /files downloads."""
import asyncio

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from harness.api.auth import get_current_user
from harness.api.routes import files as files_route
from harness.billing import entitlements
from harness.db import billing as billing_db
from harness.db.billing import Account
from harness.tools.builtin import data, files
from harness.tools.builtin.daily import build_daily_tools

STATEMENT = """Date,Description,Category,Amount
01/09/2026,Swiggy,Food,"₹450"
03/09/2026,Rent,Housing,"₹15,000"
05/09/2026,Uber,Transport,₹320
12/09/2026,BigBasket,Food,"₹2,100"
02/10/2026,Rent,Housing,"₹15,000"
04/10/2026,Zomato,Food,₹600
"""
MD = "# Budget\n\nRent is **₹15,000**.\n\n## Items\n- Rent\n- Food\n\n| Item | Amount |\n|---|---|\n| Rent | 15000 |\n\n# Next steps\n- Save more"


@pytest.fixture
def folder(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "SESSIONS", tmp_path)
    d = files.user_folder("7")
    (d / "statement.csv").write_text(STATEMENT)
    return d


def run(coro):
    return asyncio.run(coro)


# ------------------------------------------------------------ create_file

def test_documents_in_every_format(folder):
    tool = files.make_create_file_tool("7")
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK"), ("pptx", b"PK"), ("md", b"# Bu"), ("txt", b"Budg")):
        out = run(tool.handler(title="Budget", format=fmt, content=MD))
        assert out.ui["kind"] == "file" and out.ui["format"] == fmt
        assert (folder / out.ui["name"]).read_bytes().startswith(magic), fmt


def test_word_and_slides_keep_structure(folder):
    from docx import Document
    from pptx import Presentation
    tool = files.make_create_file_tool("7")
    doc = Document(folder / run(tool.handler(title="Budget", format="docx", content=MD)).ui["name"])
    assert any("₹15,000" in p.text for p in doc.paragraphs) and len(doc.tables) == 1
    deck = Presentation(folder / run(tool.handler(title="Budget", format="pptx", content=MD)).ui["name"])
    assert [s.shapes.title.text for s in deck.slides] == ["Budget", "Items", "Next steps"]


def test_pdf_keeps_unicode_text(folder):
    from pypdf import PdfReader
    name = run(files.make_create_file_tool("7").handler(title="Budget", format="pdf", content=MD)).ui["name"]
    assert "₹15,000" in PdfReader(folder / name).pages[0].extract_text()


def test_spreadsheet_cells_are_typed(folder):
    from openpyxl import load_workbook
    out = run(files.make_create_file_tool("7").handler(title="Budget", format="xlsx",
              sheets=[{"name": "Sep", "rows": [["Item", "Amount"], ["Rent", "15,000"], ["Food", 6200]]}]))
    ws = load_workbook(folder / out.ui["name"])["Sep"]
    assert ws["B2"].value == 15000 and ws["B3"].value == 6200 and ws["A1"].font.bold


def test_files_are_never_overwritten_and_bad_input_is_explained(folder):
    tool = files.make_create_file_tool("7")
    a = run(tool.handler(title="Notes", format="md", content="a")).ui["name"]
    b = run(tool.handler(title="Notes", format="md", content="b")).ui["name"]
    assert (a, b) == ("notes.md", "notes-2.md") and (folder / a).read_text() == "a"
    assert "Unsupported format" in run(tool.handler(title="x", format="exe", content="x"))
    assert "needs `sheets`" in run(tool.handler(title="x", format="xlsx", content="plain words"))


# ----------------------------------------------------------- analyze_data

def test_spend_by_category_reads_rupee_amounts(folder):
    out = run(data.make_analyze_data_tool("7").handler(
        action="query", file="statement.csv", group_by=["category"], aggregate=[{"column": "Amount", "op": "sum"}]))
    assert out.ui["rows"] == [["Category", "sum_Amount"], ["Housing", 30000], ["Food", 3150], ["Transport", 320]]


def test_group_by_month_and_filters(folder):
    tool = data.make_analyze_data_tool("7")
    out = run(tool.handler(action="query", file="statement.csv", group_by=["Date"], date_part="month",
                           aggregate=[{"column": "Amount", "op": "sum"}], sort_by="Date", descending=False))
    assert out.ui["rows"][1:] == [["2026-09", 17870], ["2026-10", 15600]]
    out = run(tool.handler(action="query", file="statement.csv",
                           filters=[{"column": "Category", "op": "=", "value": "Food"}, {"column": "Amount", "op": ">", "value": "500"}]))
    assert [r[1] for r in out.ui["rows"][1:]] == ["BigBasket", "Zomato"]


def test_chart_is_drawn_into_the_users_folder(folder):
    out = run(data.make_analyze_data_tool("7").handler(
        action="chart", file="statement.csv", group_by=["Category"], aggregate=[{"column": "Amount", "op": "sum"}],
        chart="pie", title="Where my money went"))
    assert out.ui["kind"] == "chart" and (folder / out.ui["name"]).read_bytes()[:4] == b"\x89PNG"


def test_data_errors_are_explained_not_raised(folder):
    tool = data.make_analyze_data_tool("7")
    assert "Available: statement.csv" in run(tool.handler(action="query", file="../8/secret.csv"))
    assert "No column 'Shop'" in run(tool.handler(action="query", file="statement.csv", group_by=["Shop"]))
    assert "needs `group_by`" in run(tool.handler(action="chart", file="statement.csv"))


# --------------------------------------------------------------- gating

def _gated(monkeypatch, plan):
    monkeypatch.setattr(entitlements, "billing_enabled", lambda: True)
    monkeypatch.setattr(billing_db, "get_account", lambda uid: Account(user_id=uid, plan=plan))
    return {t.name: t for t in entitlements.gate_tools("7", build_daily_tools("7"))}


def test_free_plan_gets_an_upgrade_card_instead(monkeypatch, folder):
    tools = _gated(monkeypatch, "free")
    out = run(tools["create_file"].handler(title="x", format="pdf", content="x"))
    assert out.ui == {"kind": "upgrade", "feature": "Creating files (PDF, Word, Excel, slides)", "plan": "plus", "plan_label": "Plus"}
    assert list(folder.glob("*.pdf")) == []                              # nothing was made
    assert tools["create_file"].parameter["required"] == ["title", "format"]   # same schema, model still sees it
    assert run(tools["convert"].handler(1, "km", "m")).ui["to"] == "1,000 m"     # ungated tools untouched


def test_plus_and_billing_off_get_the_real_tools(monkeypatch, folder):
    tools = _gated(monkeypatch, "plus")
    assert run(tools["create_file"].handler(title="x", format="md", content="x")).ui["kind"] == "file"
    monkeypatch.setattr(entitlements, "billing_enabled", lambda: False)
    real = {t.name: t for t in build_daily_tools("7")}
    assert all(t is real[t.name] for t in entitlements.gate_tools("7", list(real.values())))   # nothing stubbed


# ------------------------------------------------------------- downloads

def _client(user_id="7"):
    app = FastAPI()
    app.include_router(files_route.router)
    app.dependency_overrides[get_current_user] = lambda: {"user_id": user_id}
    return TestClient(app)


def test_download_own_file_only(folder):
    (folder / "report.pdf").write_bytes(b"%PDF-1.4 x")
    (folder / "chart.png").write_bytes(b"\x89PNGx")
    r = _client().get("/files/report.pdf")
    assert r.status_code == 200 and r.content == b"%PDF-1.4 x" and "attachment" in r.headers["content-disposition"]
    assert "inline" in _client().get("/files/chart.png").headers["content-disposition"]
    assert _client("8").get("/files/report.pdf").status_code == 404
    assert _client().get("/files/..%2F8%2Fx.pdf").status_code == 404
    assert _client().get("/files/.env").status_code == 404


def test_file_list_is_the_users_own_newest_first(folder):
    import os
    (folder / "old.pdf").write_bytes(b"%PDF")
    (folder / "chart.png").write_bytes(b"\x89PNG")
    os.utime(folder / "old.pdf", (1, 1))
    os.utime(folder / "chart.png", (2_000_000_000, 2_000_000_000))      # newest, explicitly
    rows = _client().get("/files").json()
    names = [r["name"] for r in rows]
    assert names[0] == "chart.png" and names[-1] == "old.pdf" and "statement.csv" in names
    assert {r["name"]: r["kind"] for r in rows}["chart.png"] == "image"
    assert _client("8").get("/files").json() == []
